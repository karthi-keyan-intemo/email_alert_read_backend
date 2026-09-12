import asyncio
from datetime import date, datetime, timezone
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.api.routes import email_alert as email_alert_routes
from app.core.auth import get_current_user
from app.repositories.email_alert_repository import EmailAlertFilters
from app.schemas.email_alert import EmailAlertReadRequest
from app.services.email_alert_service import EmailAlertService


client = TestClient(app)


class Permission:
    def __init__(self, name):
        self.name = name


class Role:
    permissions = [Permission("ALERT_VIEW"), Permission("ALERT_SYNC"), Permission("ALERT_EXPORT"), Permission("ALERT_EDIT")]


class AuthenticatedUser:
    roles = [Role()]




@pytest.fixture(autouse=True)
def authenticated_alert_requests():
    app.dependency_overrides[get_current_user] = lambda: AuthenticatedUser()
    yield
    app.dependency_overrides.pop(get_current_user, None)


class FakeRepository:
    def __init__(self):
        self.summary_calls = []
        self.page_calls = []

    def get_page_summary(self, **kwargs):
        self.summary_calls.append(kwargs)
        return {
            "total_alerts": 47,
            "unknown_errors": 30,
            "response_validation_errors": 17,
            "total_requests": 86,
        }

    def find_all(self, **kwargs):
        self.page_calls.append(kwargs)
        return []

    def get_analytics(self, filters):
        self.analytics_filters = filters
        return {"summary": {"total_alerts": 0}}


def test_service_passes_database_pagination_and_calculates_total_pages():
    repository = FakeRepository()
    service = EmailAlertService(email_reader=None, repository=repository)

    result = service.get_email_alerts(
        from_date=date(2026, 9, 8),
        to_date=date(2026, 9, 10),
        page=2,
        page_size=20,
    )

    assert result["total"] == 47
    assert result["total_pages"] == 3
    assert repository.page_calls[0]["offset"] == 20
    assert repository.page_calls[0]["limit"] == 20
    assert repository.page_calls[0]["from_date"] == date(2026, 9, 8)
    assert repository.page_calls[0]["to_date"] == date(2026, 9, 10)


def test_service_passes_analytics_filters_to_repository():
    repository = FakeRepository()
    service = EmailAlertService(email_reader=None, repository=repository)
    filters = EmailAlertFilters(from_date=date(2026, 9, 8), to_date=date(2026, 9, 10), source_name="ONEY")

    result = service.get_analytics(filters)

    assert result["summary"]["total_alerts"] == 0
    assert repository.analytics_filters == filters


@pytest.mark.parametrize(
    "query",
    ["page=0", "page_size=0", "page_size=101"],
)
def test_get_alerts_rejects_invalid_pagination(query):
    response = client.get(f"/api/email-alerts?{query}")

    assert response.status_code == 422


def test_get_and_export_reject_reversed_date_ranges():
    query = "from_date=2026-09-10&to_date=2026-09-08"

    assert client.get(f"/api/email-alerts?{query}").status_code == 400
    assert client.get(f"/api/email-alerts/export?{query}").status_code == 400


def test_analytics_rejects_reversed_date_ranges():
    response = client.get("/api/email-alerts/analytics?from_date=2026-09-10&to_date=2026-09-08")

    assert response.status_code == 400


def test_sync_request_requires_a_valid_date_range():
    valid = EmailAlertReadRequest(from_date=date(2026, 9, 10), to_date=date(2026, 9, 10))
    assert valid.from_date == valid.to_date

    with pytest.raises(ValueError):
        EmailAlertReadRequest(from_date=date(2026, 9, 11), to_date=date(2026, 9, 10))


def test_alerts_reject_unsupported_sort_field():
    response = client.get("/api/email-alerts?sort_by=not_a_column")

    assert response.status_code == 422


def test_export_returns_one_excel_row_per_alert(monkeypatch):
    class Request:
        def __init__(self, request_id):
            self.request_id = request_id

    class Alert:
        id = "alert-1"
        email_subject = "Subject"
        error_type = "UNKNOWN"
        environment = "prod"
        source_name = "rates"
        error_message = "first line\nsecond line"
        azure_task = None
        alert_timestamp = datetime(2026, 9, 8, 12, 30, tzinfo=timezone.utc)
        email_received_at = datetime(2026, 9, 8, 13, 30, tzinfo=timezone.utc)
        sender_email = "sender@example.com"
        requests = [Request("request-1"), Request("request-2")]

    class Repository:
        def __init__(self, db):
            pass

        def find_all_for_export(self, **kwargs):
            return [Alert()]

    monkeypatch.setattr(email_alert_routes, "EmailAlertRepository", Repository)
    response = email_alert_routes.export_email_alerts(
        from_date=date(2026, 9, 8),
        to_date=date(2026, 9, 10),
        db=None,
    )
    async def read_body():
        return b"".join([chunk async for chunk in response.body_iterator])

    body = asyncio.run(read_body())

    assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "email_alerts_2026-09-08_to_2026-09-10.xlsx" in response.headers["content-disposition"]
    workbook = load_workbook(BytesIO(body))
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[0] == ("Alert ID", "Environment", "Source Name", "Error Message", "Azure Task", "Request IDs")
    assert rows[1][0] == "alert-1"
    assert rows[1][3] == "first line\nsecond line"
    assert rows[1][4] is None
    assert rows[1][5] == "request-1, request-2"
    assert len(rows) == 2